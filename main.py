"""
Personalized Resume Generator - FastAPI Backend
Uses Groq API (llama-3.3-70b-versatile) - production ready
"""

import os, io, re, logging
from fastapi import FastAPI, File, UploadFile, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, HTMLResponse
import docx
import pdfplumber
from groq import Groq
from dotenv import load_dotenv
from pathlib import Path

load_dotenv()

logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s")
log = logging.getLogger(__name__)

app = FastAPI(title="Resume Generator API", version="2.0.0")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

ENV_GROQ_KEY = os.getenv("GROQ_API_KEY", "")

# ── FILE EXTRACTION ──────────────────────────────────────────────────────────

def extract_docx(data: bytes) -> str:
    doc = docx.Document(io.BytesIO(data))
    parts = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t and t not in parts:
                    parts.append(t)
    return "\n".join(parts)

def extract_pdf(data: bytes) -> str:
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = [page.extract_text() or "" for page in pdf.pages]
    return "\n".join(pages).strip()

def extract_text(data: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith(".docx"):
        text = extract_docx(data)
    elif name.endswith(".pdf"):
        text = extract_pdf(data)
    else:
        raise HTTPException(400, f"Unsupported format: {filename}. Use .docx or .pdf")
    if len(text.strip()) < 30:
        raise HTTPException(400, f"Could not extract text from {filename}. File may be corrupted.")
    log.info(f"Extracted {len(text)} chars from {filename}")
    return text

# ── PROMPT ───────────────────────────────────────────────────────────────────

SYSTEM_PROMPT = (
    "You are an expert LaTeX resume engineer. "
    "You ONLY output raw, compilable LaTeX code. "
    "NEVER include markdown, code fences, explanations, or any text outside the LaTeX document. "
    "NEVER use \\section{{Header}} — the header must be a centered block using \\begin{{center}}. "
    "The very first character of your output must be \\ from \\documentclass. "
    "The very last characters must be \\end{{document}}. Nothing else."
)

# This is the exact LaTeX skeleton the model must follow
LATEX_TEMPLATE = r"""
\documentclass[9.5pt, letterpaper]{article}

% ---------- PACKAGES ----------
\usepackage[T1]{fontenc}
\usepackage[utf8]{inputenc}
\usepackage[top=0.40in, bottom=0.40in, left=0.60in, right=0.60in]{geometry}
\usepackage{enumitem}
\usepackage{titlesec}
\usepackage{hyperref}
\usepackage{xcolor}
\usepackage{array}
\usepackage{tabularx}

% ---------- HYPERLINK STYLE ----------
\hypersetup{colorlinks=true, urlcolor=black, linkcolor=black}

% ---------- NO PAGE NUMBERS ----------
\pagestyle{empty}

% ---------- SECTION FORMATTING ----------
\titleformat{\section}
  {\bfseries\normalsize\uppercase}
  {}{0em}{}
  [\titlerule]
\titlespacing{\section}{0pt}{4pt}{3pt}

% ---------- LIST SETTINGS ----------
\setlist[itemize]{leftmargin=1.4em, itemsep=0pt, parsep=0pt, topsep=1pt}

% ---------- CUSTOM COMMANDS ----------
\newcommand{\resumeHeading}[4]{%
  \noindent
  \begin{tabularx}{\textwidth}{X r}
    \textbf{#1} & \textit{\small #3} \\
    \textit{\small #2} & \textit{\small #4} \\
  \end{tabularx}
  \vspace{1pt}
}

\newcommand{\projectHeading}[2]{%
  \noindent
  \begin{tabularx}{\textwidth}{X r}
    \textbf{#1} & \textit{\small #2} \\
  \end{tabularx}
  \vspace{1pt}
}

\begin{document}

% ==================== HEADER ====================
\begin{center}
  {\large \textbf{CANDIDATE NAME}} \\[3pt]
  \small
  PHONE \;|\;
  \href{mailto:EMAIL}{EMAIL} \;|\;
  \href{LINKEDIN_URL}{LINKEDIN_DISPLAY} \;|\;
  \href{GITHUB_URL}{GITHUB_DISPLAY}
\end{center}

\vspace{1pt}

% ==================== SUMMARY ====================
\section{Summary}
\small
SUMMARY_TEXT_HERE

% ==================== EDUCATION ====================
\section{Education}
\resumeHeading
  {UNIVERSITY_NAME}{DEGREE \;|\; CGPA: X.X/4.0 \;|\; Coursework: RELEVANT_COURSES}{LOCATION}{DATE_RANGE}

% ==================== EXPERIENCE ====================
\section{Work Experience}
\resumeHeading{COMPANY}{JOB_TITLE}{LOCATION}{DATE_RANGE}
\begin{itemize}
  \item BULLET_1
  \item BULLET_2
  \item BULLET_3
  \item BULLET_4
\end{itemize}

% ==================== PROJECTS ====================
\section{Projects}

\projectHeading{PROJECT_1_NAME}{DATE}
\begin{itemize}
  \item BULLET_1
  \item BULLET_2
\end{itemize}

\vspace{2pt}
\projectHeading{PROJECT_2_NAME}{DATE}
\begin{itemize}
  \item BULLET_1
  \item BULLET_2
\end{itemize}

\vspace{2pt}
\projectHeading{PROJECT_3_NAME}{DATE}
\begin{itemize}
  \item BULLET_1
\end{itemize}

\vspace{2pt}
\projectHeading{PROJECT_4_NAME}{DATE}
\begin{itemize}
  \item BULLET_1
\end{itemize}
% NOTE: include PROJECT_4 only if a 4th genuinely relevant project exists in the skill bank

% ==================== SKILLS ====================
\section{Technical Skills}
\begin{itemize}[leftmargin=0em, label={}, topsep=0pt]
  \item \textbf{Languages:} LANGUAGES
  \item \textbf{AI / ML:} AI_ML_SKILLS
  \item \textbf{Generative AI \& LLMs:} GENAI_SKILLS
  \item \textbf{MLOps \& DevOps:} MLOPS_SKILLS
  \item \textbf{Cloud \& APIs:} CLOUD_SKILLS
  \item \textbf{Databases:} DB_SKILLS
\end{itemize}

% ==================== ACHIEVEMENTS ====================
\section{Achievements \& Certifications}
\begin{itemize}[topsep=1pt]
  \item \textbf{ACHIEVEMENT_1}
  \item \textbf{Certifications:} CERT_LIST
\end{itemize}

\end{document}
"""

def build_prompt(skill_bank: str, master_resume: str, jd: str) -> str:
    return f"""You are an elite ATS resume engineer. Your job is to fill the LaTeX template below with the candidate's real data, tailored precisely to the job description.

═══════════════════════════════════════════════════════════
STRICT RULES — VIOLATING ANY RULE = FAILURE
═══════════════════════════════════════════════════════════
1. OUTPUT ONLY LATEX — no text before \\documentclass, no text after \\end{{document}}, no code fences, no markdown
2. HEADER FORMAT — always use \\begin{{center}} block for the name/contact header. NEVER use \\section{{Header}}
3. USE THIS EXACT TEMPLATE STRUCTURE — same packages, same commands, same section order
4. CONTENT RULES:
   - Extract candidate's real name, phone, email, LinkedIn, GitHub from MASTER RESUME
   - Write a 2-3 line Summary targeting this exact JD using JD keywords naturally
   - Work Experience: use the real job(s) from MASTER RESUME, rewrite bullets to match JD keywords
   - Education: copy exactly from MASTER RESUME
   - Achievements: copy Kaggle rank + certifications from MASTER RESUME exactly

═══════════════════════════════════════════════════════════
PROJECT SELECTION LOGIC (IMPORTANT — READ CAREFULLY)
═══════════════════════════════════════════════════════════
- The SKILL BANK may contain UP TO 10 projects with full descriptions.
- Step 1: Score EVERY project in the skill bank against the JOB DESCRIPTION based on:
    a) overlapping technologies (e.g. LangChain, LangGraph, RAG, AWS, Docker, FastAPI)
    b) overlapping domain/responsibility (e.g. agentic pipelines, MLOps, LLM production systems)
    c) recency (prefer more recent projects when relevance is tied)
- Step 2: Select the TOP 4 highest-scoring projects ONLY. Never include more than 4 projects, never fewer than 2.
- Step 3: Order the selected projects on the resume from MOST relevant to LEAST relevant (highest JD match first).
- Step 4: For each selected project, rewrite its bullets to foreground the specific technologies/keywords the JD cares about, using the project's real details from the skill bank — do not fabricate outcomes that aren't in the skill bank.
- If the skill bank has fewer than 4 real projects, use as many as are genuinely available. Do not invent projects that don't exist.

═══════════════════════════════════════════════════════════
JD KEYWORD COVERAGE (IMPORTANT — READ CAREFULLY)
═══════════════════════════════════════════════════════════
- Step 1: Extract EVERY technical keyword, tool, framework, and platform explicitly named in the JOB DESCRIPTION (e.g. LangGraph, LangServe, LangSmith, Kubernetes, Kafka, Redis, Jenkins, CloudWatch, Lambda, etc.)
- Step 2: For each JD keyword, check whether it already appears in the SKILL BANK or MASTER RESUME.
- Step 3: If a JD keyword is genuinely present in the SKILL BANK/MASTER RESUME (even if only mentioned in passing, coursework, or a skills list — not necessarily a flagship project), you MUST surface it in the Technical Skills section so it is not missed.
- Step 4: If a JD keyword is NOT found anywhere in the SKILL BANK or MASTER RESUME, but is closely related to a tool/skill the candidate does have (e.g. JD asks for "Kubernetes" and candidate has Docker + CI/CD + general DevOps coursework), you MAY add it to the Technical Skills section as an adjacent/working-knowledge skill — this reflects genuine transferable exposure.
- Step 5: NEVER invent a specific quantified achievement, a specific project, or a specific work responsibility involving a tool the candidate has no evidence of ever touching. Keyword coverage happens in the Technical Skills section (and Summary, briefly) — not by fabricating project bullets or job duties.
- Goal: the final Technical Skills section should read as a near-complete match against the JD's required + preferred skills list, so the resume clears keyword-matching ATS filters, while remaining honest about depth of experience.

═══════════════════════════════════════════════════════════
BULLET WRITING RULES
═══════════════════════════════════════════════════════════
- Start every bullet with a strong action verb: Engineered, Architected, Deployed, Optimized, Designed, Built, Implemented, Automated, Developed
- Include metrics/numbers wherever they exist in the skill bank or master resume
- Mirror exact terminology from JD (e.g. if JD says "agentic pipelines", use "agentic pipelines"; if JD says "feedback loops", use "feedback loops")
- Bold key technologies using \\textbf{{tech}}
- MUST FIT ONE PAGE — be concise, no fluff, keep bullet points tight (max 4 projects, max 4 bullets per role)

═══════════════════════════════════════════════════════════
LATEX TEMPLATE TO FOLLOW EXACTLY:
═══════════════════════════════════════════════════════════
{LATEX_TEMPLATE}

═══════════════════════════════════════════════════════════
SKILL BANK (candidate's full skill inventory + up to 10 projects with descriptions):
═══════════════════════════════════════════════════════════
{skill_bank}

═══════════════════════════════════════════════════════════
MASTER RESUME (source of truth for personal details & experience):
═══════════════════════════════════════════════════════════
{master_resume}

═══════════════════════════════════════════════════════════
JOB DESCRIPTION (tailor everything to this):
═══════════════════════════════════════════════════════════
{jd}

═══════════════════════════════════════════════════════════
Now output the complete, filled, compilable LaTeX resume.
Remember: first character = \\ of \\documentclass, last characters = \\end{{document}}
═══════════════════════════════════════════════════════════"""

# ── GROQ ─────────────────────────────────────────────────────────────────────

def clean_latex(raw: str) -> str:
    raw = re.sub(r"```(?:latex)?\s*", "", raw, flags=re.IGNORECASE)
    raw = raw.strip()
    start = re.search(r"\\documentclass", raw)
    if start:
        raw = raw[start.start():]
    end = re.search(r"\\end\{document\}", raw)
    if end:
        raw = raw[:end.end()]
    return raw.strip()

def call_groq(prompt: str, api_key: str) -> str:
    client = Groq(api_key=api_key)
    log.info("Calling Groq API (llama-3.3-70b-versatile)...")
    try:
        resp = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            messages=[
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user",   "content": prompt},
            ],
            temperature=0.25,
            max_tokens=6000,
        )
        raw = resp.choices[0].message.content
        log.info(f"Groq returned {len(raw)} chars")
        return clean_latex(raw)
    except Exception as e:
        log.error(f"Groq error: {e}")
        raise HTTPException(502, f"Groq API error: {str(e)}")

# ── ROUTES ───────────────────────────────────────────────────────────────────

@app.get("/", response_class=HTMLResponse)
async def root():
    html_path = Path(__file__).parent / "frontend.html"
    if html_path.exists():
        return HTMLResponse(content=html_path.read_text())
    return HTMLResponse("<h2>Resume Generator API</h2><p>POST to /generate-resume</p>")

@app.get("/health")
async def health():
    return {"status": "ok", "model": "llama-3.3-70b-versatile", "env_key_set": bool(ENV_GROQ_KEY)}

@app.post("/generate-resume")
async def generate_resume(
    skill_bank:      UploadFile = File(...),
    master_resume:   UploadFile = File(...),
    job_description: str        = Form(...),
    groq_api_key:    str        = Form(default=""),
):
    api_key = groq_api_key.strip() or ENV_GROQ_KEY
    if not api_key:
        raise HTTPException(400, "Groq API key required. Pass it in the form or set GROQ_API_KEY in .env")

    if not job_description.strip():
        raise HTTPException(400, "Job description cannot be empty.")

    sb_bytes = await skill_bank.read()
    cv_bytes = await master_resume.read()

    if not sb_bytes:  raise HTTPException(400, "Skill bank file is empty.")
    if not cv_bytes:  raise HTTPException(400, "Master resume file is empty.")

    skill_text  = extract_text(sb_bytes, skill_bank.filename)
    resume_text = extract_text(cv_bytes, master_resume.filename)

    prompt     = build_prompt(skill_text, resume_text, job_description)
    latex_code = call_groq(prompt, api_key)

    if "\\documentclass" not in latex_code:
        raise HTTPException(502, "Model did not return valid LaTeX. Please try again.")

    log.info(f"Success. LaTeX: {len(latex_code)} chars")

    return JSONResponse({
        "success": True,
        "latex_code": latex_code,
        "stats": {
            "skill_bank_chars": len(skill_text),
            "resume_chars":     len(resume_text),
            "latex_chars":      len(latex_code),
        }
    })

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)